"""Extractor suite tests (PYE-02, D-04) — fixtures generated inline.

Covers the dispatcher gates (type allowlist, size cap), TXT/CSV/DOCX/XLSX
parsing, PDF per-page digital routing (pypdf-generated text-layer PDF), and
the OCR path (skips without a tesseract binary — Docker image only,
RESEARCH Pitfall 6; the dev host has none).
"""

import shutil
from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, StreamObject

from app.documents.extractors import extract_document


async def test_extract_txt_roundtrip(tmp_path: Path) -> None:
    """(a) Dispatcher returns TXT content as-is."""
    p = tmp_path / "note.txt"
    p.write_text("hello academio", encoding="utf-8")
    res = await extract_document(str(p))
    assert res.text == "hello academio"
    assert res.chars == len("hello academio")
    assert res.pages == 1


async def test_extract_csv_rows_joined(tmp_path: Path) -> None:
    """(b) CSV rows are joined with ' | ' (office.py contract)."""
    p = tmp_path / "grades.csv"
    p.write_text("name,grade\nada,95\n", encoding="utf-8")
    res = await extract_document(str(p))
    assert "name | grade" in res.text
    assert "ada | 95" in res.text


async def test_extract_unsupported_type(tmp_path: Path) -> None:
    """(c) Unknown extension -> ValueError (-> 400 at the route)."""
    p = tmp_path / "file.xyz"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError, match="unsupported document type"):
        await extract_document(str(p))


async def test_extract_oversized_file_rejected(tmp_path: Path) -> None:
    """(d) > AI_MAX_DOC_MB (50MB) -> ValueError BEFORE parsing (T-03-05-03)."""
    p = tmp_path / "huge.pdf"
    with open(p, "wb") as f:
        f.truncate(51 * 1024 * 1024)  # sparse file — apparent size 51MB
    with pytest.raises(ValueError, match="size limit"):
        await extract_document(str(p))


def _write_text_pdf(path: str) -> None:
    """Build a minimal 1-page PDF with a real text layer via pypdf."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    page[NameObject("/Resources")] = resources
    stream = StreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (Hello Academio digital PDF text.) Tj ET")
    page[NameObject("/Contents")] = stream
    with open(path, "wb") as f:
        writer.write(f)


async def test_extract_pdf_digital_routing(tmp_path: Path) -> None:
    """(e) Text-layer page (>20 chars) -> digital path, no OCR attempted."""
    p = tmp_path / "digital.pdf"
    _write_text_pdf(str(p))
    # Sanity: pypdf can actually read the text back off the fixture.
    assert (PdfReader(str(p)).pages[0].extract_text() or "").strip() == (
        "Hello Academio digital PDF text."
    )
    res = await extract_document(str(p))
    assert res.pages == 1
    assert res.ocr_pages == 0  # routed digital, never touched tesseract
    assert "Hello Academio digital PDF text." in res.text


@pytest.mark.skipif(
    shutil.which("tesseract") is None,
    reason="tesseract binary required (Docker image only — RESEARCH Pitfall 6)",
)
async def test_extract_image_ocr(tmp_path: Path) -> None:
    """(f) OCR path: synthetic image with rendered text -> non-empty text."""
    img = Image.new("RGB", (800, 200), "white")
    ImageDraw.Draw(img).text((20, 40), "Academio OCR test", fill="black")
    p = tmp_path / "ocr.png"
    img.save(p)
    res = await extract_document(str(p))
    assert res.ocr_pages == 1
    assert res.text.strip() != ""


async def test_extract_docx(tmp_path: Path) -> None:
    """DOCX paragraph text via python-docx writer (extra coverage)."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Academio docx paragraph")
    p = tmp_path / "doc.docx"
    doc.save(str(p))
    res = await extract_document(str(p))
    assert "Academio docx paragraph" in res.text


async def test_extract_xlsx(tmp_path: Path) -> None:
    """XLSX cells joined with ' | ' via openpyxl writer (extra coverage)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["name", "score"])
    ws.append(["ada", 95])
    p = tmp_path / "sheet.xlsx"
    wb.save(str(p))
    res = await extract_document(str(p))
    assert "name | score" in res.text
    assert "ada | 95" in res.text
